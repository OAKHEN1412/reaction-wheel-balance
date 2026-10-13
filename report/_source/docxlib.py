# Helpers for Thai thesis-style .docx (TH Sarabun New 16pt, A4, 1.5"/1" margins)
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_COLOR_INDEX, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from xml.sax.saxutils import escape

FONT = 'TH Sarabun New'
SIZE = 16
TEXT_W_CM = 21.0 - 2.54 * 2.5  # 14.65 cm
MNS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" ' \
      'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline',
             'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing', 'w',
             'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign',
             'rtl', 'cs', 'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath']


def sort_rpr(rPr):
    kids = list(rPr)
    def key(el):
        tag = el.tag.split('}')[-1]
        return RPR_ORDER.index(tag) if tag in RPR_ORDER else len(RPR_ORDER)
    kids.sort(key=key)
    for k in kids:
        rPr.remove(k)
    for k in kids:
        rPr.append(k)


def _fonts(rPr, size=SIZE, bold=False):
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    for tag in ('w:b', 'w:bCs'):
        el = rPr.find(qn(tag))
        if bold and el is None:
            rPr.append(OxmlElement(tag))
    for tag in ('w:sz', 'w:szCs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rPr.append(el)
        el.set(qn('w:val'), str(int(size * 2)))
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang'); rPr.append(lang)
    lang.set(qn('w:val'), 'en-US'); lang.set(qn('w:bidi'), 'th-TH')
    sort_rpr(rPr)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Inches(1.5); sec.left_margin = Inches(1.5)
    sec.bottom_margin = Inches(1.0); sec.right_margin = Inches(1.0)
    sec.header_distance = Inches(0.5); sec.footer_distance = Inches(0.5)
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(SIZE)
    _fonts(st.element.get_or_add_rPr())
    pf = st.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    # Table Grid style fonts
    for s in doc.styles:
        try:
            _fonts(s.element.get_or_add_rPr())
        except Exception:
            pass
    return doc


class M:
    """Inline math built from a tiny spec. Items: str (plain math text), ('sub', base, sub),
    ('dot', base), ('frac', [num items], [den items]), ('int', [items]), ('sup', base, sup)."""
    def __init__(self, *items):
        self.items = items


MATH_SZ = [24]


def _mr(t, plain=False):
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ''
    sz = MATH_SZ[0]
    return (f'<m:r>{sty}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
            f'<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr><m:t xml:space="preserve">{escape(t)}</m:t></m:r>')


def _mitems(items):
    out = []
    for it in items:
        if isinstance(it, str):
            # digits/operators upright, letters italic handled by Word defaults
            out.append(_mr(it, plain=not any(c.isalpha() for c in it)))
        elif it[0] == 'txt':
            out.append(_mr(it[1], plain=True))
        elif it[0] == 'sub':
            b = _mitems([it[1]] if isinstance(it[1], str) else it[1])
            s = _mitems([it[2]] if isinstance(it[2], str) else it[2])
            out.append(f'<m:sSub><m:e>{b}</m:e><m:sub>{s}</m:sub></m:sSub>')
        elif it[0] == 'sup':
            b = _mitems([it[1]] if isinstance(it[1], str) else it[1])
            s = _mitems([it[2]] if isinstance(it[2], str) else it[2])
            out.append(f'<m:sSup><m:e>{b}</m:e><m:sup>{s}</m:sup></m:sSup>')
        elif it[0] == 'dot':
            b = _mitems([it[1]] if isinstance(it[1], str) else it[1])
            out.append(f'<m:acc><m:accPr><m:chr m:val="̇"/></m:accPr><m:e>{b}</m:e></m:acc>')
        elif it[0] == 'ddot':
            b = _mitems([it[1]] if isinstance(it[1], str) else it[1])
            out.append(f'<m:acc><m:accPr><m:chr m:val="\u0308"/></m:accPr><m:e>{b}</m:e></m:acc>')
        elif it[0] == 'frac':
            out.append(f'<m:f><m:num>{_mitems(it[1])}</m:num><m:den>{_mitems(it[2])}</m:den></m:f>')
        elif it[0] == 'int':
            out.append('<m:nary><m:naryPr><m:chr m:val="∫"/><m:limLoc m:val="subSup"/>'
                       '<m:subHide m:val="1"/><m:supHide m:val="1"/></m:naryPr><m:sub/><m:sup/>'
                       f'<m:e>{_mitems(it[1])}</m:e></m:nary>')
        elif it[0] == 'paren':
            out.append(f'<m:d><m:dPr><m:ctrlPr/></m:dPr><m:e>{_mitems(it[1])}</m:e></m:d>')
    return ''.join(out)


def _omath(m, size=12):
    MATH_SZ[0] = int(size * 2)
    try:
        return parse_xml(f'<m:oMath {MNS}>{_mitems(m.items)}</m:oMath>')
    finally:
        MATH_SZ[0] = 24


import re as _re
from pythainlp.tokenize import word_tokenize as _wt
_THAI = _re.compile('[\u0E00-\u0E7F]')
ZWSP = '\u200b'


_NO_BREAK_AFTER = '([{"“‘'
_NO_BREAK_BEFORE = ')]}"”’,.:;'


def thai_breaks(text):
    """Insert zero-width spaces between Thai words so Word can wrap Thai lines without a Thai proofing pack."""
    if not _THAI.search(text):
        return text
    toks = _wt(text, engine='newmm', keep_whitespace=True)
    out = [toks[0]] if toks else []
    for a, b in zip(toks, toks[1:]):
        # never break just inside a bracket or quote, e.g. "SolidWorks (|โปรแกรม..."
        if (a.strip() and b.strip() and (_THAI.search(a[-1]) or _THAI.search(b[0]))
                and a[-1] not in _NO_BREAK_AFTER and b[0] not in _NO_BREAK_BEFORE):
            out.append(ZWSP)
        out.append(b)
    return ''.join(out)


def add_runs(p, parts, size=SIZE, bold=False, highlight=False):
    for part in parts:
        if isinstance(part, M):
            p._p.append(_omath(part))
            continue
        if isinstance(part, tuple):  # (text, dict(style))
            text, sty = part
        else:
            text, sty = part, {}
        r = p.add_run(thai_breaks(text))
        _fonts(r._element.get_or_add_rPr(), size=sty.get('size', size), bold=sty.get('bold', bold))
        if sty.get('bold', bold):
            r.bold = True
        if sty.get('italic'):
            r.italic = True
        if sty.get('sub'):
            r.font.subscript = True
        if sty.get('sup'):
            r.font.superscript = True
        if sty.get('hl', highlight):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def para(doc, *parts, align='justify', first=0.5, left=0.0, hanging=None, bold=False, size=SIZE,
         keep_next=False, highlight=False, space_before=0, space_after=0, tabs=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {'justify': WD_ALIGN_PARAGRAPH.THAI_JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT,
                    'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf.left_indent = Inches(left)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    else:
        pf.first_line_indent = Inches(first)
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after)
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if tabs:
        for pos, kind in tabs:
            pf.tab_stops.add_tab_stop(Cm(pos), {'l': WD_TAB_ALIGNMENT.LEFT, 'c': WD_TAB_ALIGNMENT.CENTER,
                                                'r': WD_TAB_ALIGNMENT.RIGHT}[kind])
    add_runs(p, parts, size=size, bold=bold, highlight=highlight)
    return p


def blank(doc):
    return para(doc, '', first=0)


def chapter(doc, num, title):
    para(doc, f'บทที่ {num}', align='center', first=0, bold=True, keep_next=True)
    para(doc, title, align='center', first=0, bold=True, keep_next=True)
    blank(doc)


def section(doc, num, title):
    blank(doc)
    para(doc, f'{num}\t{title}', align='left', first=0, bold=True, keep_next=True, tabs=[(1.27, 'l')])


def subsection(doc, num, title):
    para(doc, f'{num} {title}', align='left', first=0.5, keep_next=True)


def item_head(doc, text):
    para(doc, text, align='left', first=0.5, keep_next=True)


def bullets(doc, items, left=0.75, hang=0.25, highlight=False):
    for it in items:
        parts = it if isinstance(it, list) else [it]
        para(doc, '-\t', *parts, left=left, hanging=hang, tabs=[(Inches(left).cm, 'l')], highlight=highlight)


def numbered(doc, items, left=0.75, hang=0.3, fmt='{}) '):
    for i, it in enumerate(items, 1):
        parts = it if isinstance(it, list) else [it]
        para(doc, fmt.format(i) + '\t', *parts, left=left, hanging=hang, tabs=[(Inches(left).cm, 'l')])


def equation(doc, m, number):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [TEXT_W_CM - 2.2, 2.2]
    for c, w in zip(t.rows[0].cells, widths):
        c.width = Cm(w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    c0, c1 = t.rows[0].cells
    p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    om = _omath(m, size=15)
    mp = parse_xml(f'<m:oMathPara {MNS}><m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr></m:oMathPara>')
    mp.append(om)
    p0._p.append(mp)
    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(p1, [f'({number})'])
    for p in (p0, p1):
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Inches(0)
    return t


def where(doc, rows):
    para(doc, 'โดยที่', align='left', first=0.5, keep_next=True)
    for sym, desc in rows:
        parts = [sym] if isinstance(sym, M) else [sym]
        para(doc, *parts, '\tคือ ', desc, align='left', first=0, left=0.5, tabs=[(Inches(1.4).cm, 'l')])


def figure(doc, path, num, caption, width_cm=14.0):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = para(doc, (f'รูปที่ {num}', {'bold': True}), f' {caption}', align='center', first=0)
    c.paragraph_format.space_after = Pt(6)
    return c


def figure_placeholder(doc, num, caption, note):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]; cell.width = Cm(12)
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr(); h = OxmlElement('w:trHeight'); h.set(qn('w:val'), '3400'); trPr.append(h)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, [note], highlight=True)
    c = para(doc, (f'รูปที่ {num}', {'bold': True}), f' {caption}', align='center', first=0)
    c.paragraph_format.space_before = Pt(6); c.paragraph_format.space_after = Pt(6)


def table(doc, num, caption, header, rows, widths_cm, size=14, align=None, highlight_rows=(), header_rows=1):
    cap = para(doc, (f'ตารางที่ {num}', {'bold': True}), f' {caption}', align='left', first=0, keep_next=True)
    cap.paragraph_format.space_before = Pt(6); cap.paragraph_format.space_after = Pt(3)
    hdrs = header if header and isinstance(header[0], list) else [header]
    t = doc.add_table(rows=len(hdrs) + len(rows), cols=len(widths_cm))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    align = align or ['c'] * len(widths_cm)
    allrows = [(h, True, False) for h in hdrs] + [(r, False, i in highlight_rows) for i, r in enumerate(rows)]
    for ri, (row, is_h, hl) in enumerate(allrows):
        tr = t.rows[ri]
        if is_h:
            trPr = tr._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
        ci = 0
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Cm(widths_cm[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            a = 'c' if is_h else align[ci]
            p.alignment = {'c': WD_ALIGN_PARAGRAPH.CENTER, 'l': WD_ALIGN_PARAGRAPH.LEFT,
                           'j': WD_ALIGN_PARAGRAPH.THAI_JUSTIFY}[a]
            parts = val if isinstance(val, list) else [val]
            add_runs(p, parts, size=size, bold=is_h, highlight=hl)
            if is_h:
                tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'E7E6E6'); tcPr.append(shd)
    # fixed layout
    tblPr = t._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    after = para(doc, '', first=0)
    return t


def merge_header(t, r, c0, c1):
    a = t.cell(r, c0); b = t.cell(r, c1)
    m = a.merge(b)
    # remove duplicate text left by merge
    ps = m.paragraphs
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    return m


def page_break(doc):
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
